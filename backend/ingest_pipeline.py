import hashlib
import json
import hmac
import os
import secrets
import shutil
import subprocess
import tempfile
import uuid
from datetime import datetime, timezone
from pathlib import Path

import docx
import faiss
import fitz
import numpy as np
import pandas as pd
import xlrd
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sentence_transformers import SentenceTransformer
from sqlalchemy import Column, Integer, MetaData, String, Table, Text, create_engine, text

from nextcloud_storage import delete_file, is_nextcloud_enabled, upload_file

try:
    import pytesseract
    from PIL import Image, ImageEnhance, ImageFilter, ImageOps
except ImportError:
    pytesseract = None
    Image = None
    ImageEnhance = None
    ImageFilter = None
    ImageOps = None


BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data_files"
INDEX_DIR = BASE_DIR / "faiss_index"
DB_PATH = BASE_DIR / "metadata.db"
SQLITE_DB = f"sqlite:///{DB_PATH}"

EMBEDDING_MODEL = os.getenv("EMBEDDING_MODEL", "sentence-transformers/all-MiniLM-L6-v2")
CHUNK_SIZE_CHARS = int(os.getenv("CHUNK_SIZE_CHARS", "1200"))
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", "200"))
MAX_FILE_SIZE_BYTES = int(os.getenv("MAX_FILE_SIZE_BYTES", str(200 * 1024 * 1024)))
SUPPORTED_EXTENSIONS = {"pdf", "docx", "doc", "xlsx", "xls"}
OLE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")
ZIP_SIGNATURE = b"PK\x03\x04"
OCR_LANGUAGE = os.getenv("OCR_LANGUAGE", "eng")
OCR_DPI_SCALE = float(os.getenv("OCR_DPI_SCALE", "2"))
OCR_THRESHOLD = int(os.getenv("OCR_THRESHOLD", "180"))
OCR_CONTRAST = float(os.getenv("OCR_CONTRAST", "1.8"))
OCR_PAGE_SEGMENT_MODES = [mode.strip() for mode in os.getenv("OCR_PSM_MODES", "6,4,11").split(",") if mode.strip()]
ACTIVE_DIRECTORY_ENABLED = os.getenv("AD_ENABLED", "false").strip().lower() in {"1", "true", "yes", "on"}
ACTIVE_DIRECTORY_DOMAIN = os.getenv("AD_DOMAIN", "").strip()
ACTIVE_DIRECTORY_SERVER = os.getenv("AD_SERVER", "").strip()
ACTIVE_DIRECTORY_DEFAULT_ROLE = os.getenv("AD_DEFAULT_ROLE", "user").strip().lower() or "user"
ACTIVE_DIRECTORY_AUTO_PROVISION = os.getenv("AD_AUTO_PROVISION", "true").strip().lower() in {"1", "true", "yes", "on"}

DATA_DIR.mkdir(parents=True, exist_ok=True)
INDEX_DIR.mkdir(parents=True, exist_ok=True)

engine = create_engine(SQLITE_DB, connect_args={"check_same_thread": False})
metadata = MetaData()

documents = Table(
    "documents",
    metadata,
    Column("id", String, primary_key=True),
    Column("filename", String, nullable=False),
    Column("stored_path", String, nullable=False),
    Column("file_hash", String, unique=True, nullable=False),
    Column("file_size", Integer, nullable=False),
    Column("file_type", String, nullable=False),
    Column("page_count", Integer, nullable=False, default=0),
    Column("chunk_count", Integer, nullable=False, default=0),
    Column("owner_id", String, nullable=True),
    Column("created_at", String, nullable=False),
)

collections = Table(
    "collections",
    metadata,
    Column("id", String, primary_key=True),
    Column("name", String, nullable=False, unique=True),
    Column("owner_id", String, nullable=True),
    Column("created_at", String, nullable=False),
)

collection_documents = Table(
    "collection_documents",
    metadata,
    Column("collection_id", String, nullable=False),
    Column("doc_id", String, nullable=False),
)

chunks = Table(
    "chunks",
    metadata,
    Column("id", String, primary_key=True),
    Column("doc_id", String, nullable=False),
    Column("chunk_index", Integer, nullable=False),
    Column("text", Text, nullable=False),
    Column("page", Integer, nullable=True),
    Column("char_start", Integer, nullable=False),
    Column("char_end", Integer, nullable=False),
    Column("section_label", String, nullable=True),
    Column("sheet_name", String, nullable=True),
    Column("cell_range", String, nullable=True),
    Column("source_ref", String, nullable=True),
)

chat_messages = Table(
    "chat_messages",
    metadata,
    Column("id", String, primary_key=True),
    Column("doc_id", String, nullable=False),
    Column("role", String, nullable=False),
    Column("message", Text, nullable=False),
    Column("user_id", String, nullable=True),
    Column("model", String, nullable=True),
    Column("citations_json", Text, nullable=True),
    Column("created_at", String, nullable=False),
)

collection_chat_messages = Table(
    "collection_chat_messages",
    metadata,
    Column("id", String, primary_key=True),
    Column("collection_id", String, nullable=False),
    Column("role", String, nullable=False),
    Column("message", Text, nullable=False),
    Column("user_id", String, nullable=True),
    Column("model", String, nullable=True),
    Column("citations_json", Text, nullable=True),
    Column("created_at", String, nullable=False),
)

metadata.create_all(engine)

users = Table(
    "users",
    metadata,
    Column("id", String, primary_key=True),
    Column("username", String, nullable=False, unique=True),
    Column("full_name", String, nullable=False),
    Column("email", String, nullable=False, unique=True),
    Column("password_hash", Text, nullable=False),
    Column("role", String, nullable=False, default="user"),
    Column("is_active", Integer, nullable=False, default=1),
    Column("auth_source", String, nullable=False, default="local"),
    Column("created_at", String, nullable=False),
)

sessions = Table(
    "sessions",
    metadata,
    Column("token", String, primary_key=True),
    Column("user_id", String, nullable=False),
    Column("created_at", String, nullable=False),
)

metadata.create_all(engine)

_embedder = None


def ensure_schema_columns():
    # Add newer columns lazily so older SQLite files keep working after app upgrades.
    with engine.begin() as conn:
        table_columns = {
            "chunks": {
                row[1]
                for row in conn.execute(text("PRAGMA table_info(chunks)")).fetchall()
            },
            "documents": {
                row[1]
                for row in conn.execute(text("PRAGMA table_info(documents)")).fetchall()
            },
            "collections": {
                row[1]
                for row in conn.execute(text("PRAGMA table_info(collections)")).fetchall()
            },
            "chat_messages": {
                row[1]
                for row in conn.execute(text("PRAGMA table_info(chat_messages)")).fetchall()
            },
            "collection_chat_messages": {
                row[1]
                for row in conn.execute(text("PRAGMA table_info(collection_chat_messages)")).fetchall()
            },
        }
        optional_columns = {
            "chunks": {
                "section_label": "ALTER TABLE chunks ADD COLUMN section_label TEXT",
                "sheet_name": "ALTER TABLE chunks ADD COLUMN sheet_name TEXT",
                "cell_range": "ALTER TABLE chunks ADD COLUMN cell_range TEXT",
                "source_ref": "ALTER TABLE chunks ADD COLUMN source_ref TEXT",
            },
            "documents": {
                "owner_id": "ALTER TABLE documents ADD COLUMN owner_id TEXT",
            },
            "collections": {
                "owner_id": "ALTER TABLE collections ADD COLUMN owner_id TEXT",
            },
            "chat_messages": {
                "user_id": "ALTER TABLE chat_messages ADD COLUMN user_id TEXT",
            },
            "collection_chat_messages": {
                "user_id": "ALTER TABLE collection_chat_messages ADD COLUMN user_id TEXT",
            },
        }
        for table_name, ddl_map in optional_columns.items():
            for column_name, ddl in ddl_map.items():
                if column_name not in table_columns[table_name]:
                    conn.execute(text(ddl))

        conn.execute(text("UPDATE documents SET owner_id = 'admin' WHERE owner_id IS NULL"))
        conn.execute(text("UPDATE collections SET owner_id = 'admin' WHERE owner_id IS NULL"))
        conn.execute(text("UPDATE chat_messages SET user_id = 'admin' WHERE user_id IS NULL"))
        conn.execute(text("UPDATE collection_chat_messages SET user_id = 'admin' WHERE user_id IS NULL"))


ensure_schema_columns()


def get_embedder():
    global _embedder
    if _embedder is None:
        _embedder = SentenceTransformer(EMBEDDING_MODEL)
    return _embedder


def _utc_now_iso():
    return datetime.now(timezone.utc).isoformat()


def _normalize_identity(value):
    return " ".join((value or "").strip().split())


def _normalize_email(value):
    return _normalize_identity(value).lower()


def _hash_password(password, salt=None):
    salt_value = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt_value.encode("utf-8"),
        120000,
    ).hex()
    return f"{salt_value}${digest}"


def _verify_password(password, password_hash):
    try:
        salt_value, stored_digest = password_hash.split("$", 1)
    except ValueError:
        return False
    candidate = _hash_password(password, salt=salt_value).split("$", 1)[1]
    return hmac.compare_digest(candidate, stored_digest)


def _display_name_from_identifier(identifier):
    cleaned = _normalize_identity(identifier)
    local_part = cleaned.split("@", 1)[0]
    words = [piece for piece in local_part.replace(".", " ").replace("_", " ").split() if piece]
    if not words:
        return cleaned
    return " ".join(word.capitalize() for word in words)


def is_active_directory_enabled():
    return ACTIVE_DIRECTORY_ENABLED and bool(ACTIVE_DIRECTORY_DOMAIN or ACTIVE_DIRECTORY_SERVER)


def authenticate_active_directory(identifier, password):
    # AD login is optional; this path only runs when real domain settings are present in the environment.
    normalized_identifier = _normalize_identity(identifier)
    if not is_active_directory_enabled() or not normalized_identifier or not password:
        return None

    principal_context_args = [ACTIVE_DIRECTORY_DOMAIN or "", ACTIVE_DIRECTORY_SERVER or ""]
    principal_context_args = [argument for argument in principal_context_args if argument]
    constructor_args = ", ".join([f"'{argument}'" for argument in principal_context_args])
    constructor = (
        f"New-Object System.DirectoryServices.AccountManagement.PrincipalContext("
        f"[System.DirectoryServices.AccountManagement.ContextType]::Domain{', ' if constructor_args else ''}{constructor_args})"
    )

    script = f"""
    try {{
      Add-Type -AssemblyName System.DirectoryServices.AccountManagement
      $context = {constructor}
      $valid = $context.ValidateCredentials('{normalized_identifier.replace("'", "''")}', '{password.replace("'", "''")}')
      if (-not $valid) {{
        Write-Output '{{"valid":false}}'
        exit 0
      }}

      $user = [System.DirectoryServices.AccountManagement.UserPrincipal]::FindByIdentity($context, '{normalized_identifier.replace("'", "''")}')
      if ($null -eq $user) {{
        Write-Output '{{"valid":true,"username":"{normalized_identifier}","email":"{normalized_identifier}","full_name":"{_display_name_from_identifier(normalized_identifier)}"}}'
        exit 0
      }}

      $result = @{{
        valid = $true
        username = if ($user.SamAccountName) {{ $user.SamAccountName }} else {{ '{normalized_identifier}' }}
        email = if ($user.EmailAddress) {{ $user.EmailAddress }} else {{ '{normalized_identifier}' }}
        full_name = if ($user.DisplayName) {{ $user.DisplayName }} else {{ '{_display_name_from_identifier(normalized_identifier)}' }}
      }}
      $result | ConvertTo-Json -Compress
    }} catch {{
      $errorResult = @{{
        valid = $false
        error = $_.Exception.Message
      }}
      $errorResult | ConvertTo-Json -Compress
      exit 0
    }}
    """

    try:
      completed = subprocess.run(
          [
              "powershell",
              "-NoProfile",
              "-NonInteractive",
              "-Command",
              script,
          ],
          check=True,
          capture_output=True,
          text=True,
          timeout=20,
      )
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
      return None

    raw_output = completed.stdout.strip()
    if not raw_output:
      return None

    try:
      payload = json.loads(raw_output.splitlines()[-1])
    except json.JSONDecodeError:
      return None

    if not payload.get("valid"):
      return None

    return {
      "username": _normalize_identity(payload.get("username") or normalized_identifier),
      "email": _normalize_email(payload.get("email") or normalized_identifier),
      "full_name": _normalize_identity(payload.get("full_name") or _display_name_from_identifier(normalized_identifier)),
      "auth_source": "active_directory",
    }


def _admin_user():
    return {
        "id": "admin",
        "username": "admin",
        "full_name": "Administrator",
        "email": "admin@local",
        "role": "admin",
        "is_active": 1,
        "auth_source": "local",
        "created_at": _utc_now_iso(),
    }


def _index_path(doc_id):
    return INDEX_DIR / f"{doc_id}.index"


def _mapping_path(doc_id):
    return INDEX_DIR / f"{doc_id}.npy"


def compute_file_hash(filepath):
    digest = hashlib.md5()
    with open(filepath, "rb") as handle:
        for chunk in iter(lambda: handle.read(8192), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _format_max_size():
    return f"{MAX_FILE_SIZE_BYTES // (1024 * 1024)}MB"


def _read_file_header(filepath, length=8):
    with open(filepath, "rb") as handle:
        return handle.read(length)


def _validate_legacy_signature(filepath, extension):
    header = _read_file_header(filepath)
    if extension in {"doc", "xls"} and not header.startswith(OLE_SIGNATURE):
        raise ValueError(
            f"The .{extension} file does not look like a valid legacy Microsoft Office document."
        )
    if extension in {"docx", "xlsx"} and not header.startswith(ZIP_SIGNATURE):
        raise ValueError(
            f"The .{extension} file appears to be invalid or has the wrong extension."
        )


def validate_upload_file(filepath, original_filename=None):
    source_path = Path(filepath)
    if not source_path.exists():
        raise FileNotFoundError(f"File not found: {filepath}")

    filename = original_filename or source_path.name
    extension = source_path.suffix.lower().lstrip(".")
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: .{extension}. Supported types are: pdf, docx, xlsx, doc, xls."
        )

    file_size = source_path.stat().st_size
    if file_size <= 0:
        raise ValueError("The uploaded file is empty.")
    if file_size > MAX_FILE_SIZE_BYTES:
        raise ValueError(
            f"File is too large. Maximum supported size is {_format_max_size()}."
        )

    _validate_legacy_signature(source_path, extension)
    return {
        "filename": filename,
        "extension": extension,
        "file_size": file_size,
    }


def get_document_by_hash(file_hash, owner_id):
    with engine.begin() as conn:
        row = conn.execute(
            text(
                """
                SELECT id, filename, stored_path, file_size, file_type, page_count, chunk_count, owner_id, created_at
                FROM documents
                WHERE file_hash = :file_hash AND owner_id = :owner_id
                """
            ),
            {"file_hash": file_hash, "owner_id": owner_id},
        ).mappings().first()
    return dict(row) if row else None


def get_document(doc_id):
    with engine.begin() as conn:
        row = conn.execute(
            text(
                """
                SELECT id, filename, stored_path, file_size, file_type, page_count, chunk_count, owner_id, created_at
                FROM documents
                WHERE id = :doc_id
                """
            ),
            {"doc_id": doc_id},
        ).mappings().first()
    return dict(row) if row else None


def get_collection(collection_id):
    with engine.begin() as conn:
        row = conn.execute(
            text(
                """
                SELECT id, name, owner_id, created_at
                FROM collections
                WHERE id = :collection_id
                """
            ),
            {"collection_id": collection_id},
        ).mappings().first()
    return dict(row) if row else None


def user_can_access_document(user, document):
    if not user or not document:
        return False
    if user.get("role") == "admin":
        return True
    return document.get("owner_id") == user.get("id")


def user_can_access_collection(user, collection):
    if not user or not collection:
        return False
    if user.get("role") == "admin":
        return True
    return collection.get("owner_id") == user.get("id")


def list_documents(current_user=None):
    if current_user and current_user.get("role") != "admin":
        where_clause = "WHERE owner_id = :owner_id"
        params = {"owner_id": current_user["id"]}
    else:
        where_clause = ""
        params = {}
    with engine.begin() as conn:
        rows = conn.execute(
            text(
                """
                SELECT id, filename, file_size, file_type, page_count, chunk_count, owner_id, created_at
                FROM documents
                {where_clause}
                ORDER BY created_at DESC
                """
                .format(where_clause=where_clause)
            ),
            params,
        ).mappings().all()
    return [dict(row) for row in rows]


def get_collection_document_ids(collection_id):
    with engine.begin() as conn:
        rows = conn.execute(
            text(
                """
                SELECT doc_id
                FROM collection_documents
                WHERE collection_id = :collection_id
                ORDER BY doc_id ASC
                """
            ),
            {"collection_id": collection_id},
        ).fetchall()
    return [row[0] for row in rows]


def list_collections(current_user=None):
    if current_user and current_user.get("role") != "admin":
        where_clause = "WHERE c.owner_id = :owner_id"
        params = {"owner_id": current_user["id"]}
    else:
        where_clause = ""
        params = {}
    with engine.begin() as conn:
        rows = conn.execute(
            text(
                """
                SELECT c.id, c.name, c.owner_id, c.created_at, COUNT(cd.doc_id) AS document_count
                FROM collections c
                LEFT JOIN collection_documents cd ON cd.collection_id = c.id
                {where_clause}
                GROUP BY c.id, c.name, c.created_at
                ORDER BY c.created_at DESC
                """
                .format(where_clause=where_clause)
            ),
            params,
        ).mappings().all()

        if not rows:
            return []

        collection_ids = [row["id"] for row in rows]
        placeholders = ", ".join(f":collection_id_{index}" for index in range(len(collection_ids)))
        params = {
            f"collection_id_{index}": collection_id
            for index, collection_id in enumerate(collection_ids)
        }
        docs = conn.execute(
            text(
                f"""
                SELECT cd.collection_id, d.id, d.filename, d.file_type
                FROM collection_documents cd
                JOIN documents d ON d.id = cd.doc_id
                WHERE cd.collection_id IN ({placeholders})
                ORDER BY d.filename ASC
                """
            ),
            params,
        ).mappings().all()

    docs_by_collection = {}
    for document in docs:
        docs_by_collection.setdefault(document["collection_id"], []).append(
            {
                "id": document["id"],
                "filename": document["filename"],
                "file_type": document["file_type"],
            }
        )

    return [
        {
            **dict(row),
            "document_ids": [doc["id"] for doc in docs_by_collection.get(row["id"], [])],
            "documents": docs_by_collection.get(row["id"], []),
        }
        for row in rows
    ]


def get_user_by_id(user_id):
    if user_id == "admin":
        return _admin_user()

    with engine.begin() as conn:
        row = conn.execute(
            text(
                """
                SELECT id, username, full_name, email, role, is_active, auth_source, created_at
                FROM users
                WHERE id = :user_id
                """
            ),
            {"user_id": user_id},
        ).mappings().first()
    return dict(row) if row else None


def list_users():
    with engine.begin() as conn:
        rows = conn.execute(
            text(
                """
                SELECT id, username, full_name, email, role, is_active, auth_source, created_at
                FROM users
                ORDER BY created_at ASC
                """
            )
        ).mappings().all()
    return [_admin_user(), *[dict(row) for row in rows]]


def create_user(full_name, email, password, username=None, role="user", is_active=1, auth_source="local"):
    normalized_full_name = _normalize_identity(full_name)
    normalized_email = _normalize_email(email)
    normalized_username = _normalize_identity(username or normalized_email)

    auth_source = (auth_source or "local").strip().lower()
    if auth_source not in {"local", "active_directory"}:
        raise ValueError("Auth source must be local or active_directory.")
    if not normalized_full_name or not normalized_email:
        raise ValueError("Full name and email are required.")
    if auth_source == "local" and not password:
        raise ValueError("Password is required for local users.")
    if role not in {"admin", "user"}:
        raise ValueError("Role must be admin or user.")

    user_id = str(uuid.uuid4())
    with engine.begin() as conn:
        existing = conn.execute(
            text(
                """
                SELECT id FROM users
                WHERE lower(email) = :email OR lower(username) = :username
                """
            ),
            {"email": normalized_email, "username": normalized_username.lower()},
        ).first()
        if existing:
            raise ValueError("A user with this email or username already exists.")

        conn.execute(
            users.insert().values(
                id=user_id,
                username=normalized_username,
                full_name=normalized_full_name,
                email=normalized_email,
                password_hash=_hash_password(password or secrets.token_urlsafe(24)),
                role=role,
                is_active=1 if is_active else 0,
                auth_source=auth_source,
                created_at=_utc_now_iso(),
            )
        )
    return get_user_by_id(user_id)


def update_user(user_id, full_name=None, email=None, password=None, role=None, is_active=None, auth_source=None):
    if user_id == "admin":
        raise ValueError("The hard admin user cannot be edited here.")

    existing_user = get_user_by_id(user_id)
    if not existing_user:
        raise FileNotFoundError("User not found.")

    updates = {
        "full_name": _normalize_identity(full_name or existing_user["full_name"]),
        "email": _normalize_email(email or existing_user["email"]),
        "role": role or existing_user["role"],
        "is_active": 1 if (existing_user["is_active"] if is_active is None else is_active) else 0,
        "auth_source": (auth_source or existing_user.get("auth_source") or "local").strip().lower(),
    }
    if updates["role"] not in {"admin", "user"}:
        raise ValueError("Role must be admin or user.")
    if updates["auth_source"] not in {"local", "active_directory"}:
        raise ValueError("Auth source must be local or active_directory.")

    with engine.begin() as conn:
        duplicate = conn.execute(
            text(
                """
                SELECT id FROM users
                WHERE lower(email) = :email AND id != :user_id
                """
            ),
            {"email": updates["email"], "user_id": user_id},
        ).first()
        if duplicate:
            raise ValueError("Another user already uses this email.")

        conn.execute(
            text(
                """
                UPDATE users
                SET full_name = :full_name,
                    email = :email,
                    role = :role,
                    is_active = :is_active,
                    auth_source = :auth_source
                WHERE id = :user_id
                """
            ),
            {**updates, "user_id": user_id},
        )
        if password:
            conn.execute(
                text(
                    """
                    UPDATE users
                    SET password_hash = :password_hash
                    WHERE id = :user_id
                    """
                ),
                {"password_hash": _hash_password(password), "user_id": user_id},
            )
    return get_user_by_id(user_id)


def delete_user_account(user_id):
    if user_id == "admin":
        raise ValueError("The hard admin user cannot be deleted.")

    user = get_user_by_id(user_id)
    if not user:
        return False

    owned_documents = list_documents(user)
    owned_collections = list_collections(user)

    for document in owned_documents:
        delete_document(document["id"])

    for collection in owned_collections:
        delete_collection(collection["id"], user)

    with engine.begin() as conn:
        conn.execute(text("DELETE FROM chat_messages WHERE user_id = :user_id"), {"user_id": user_id})
        conn.execute(
            text("DELETE FROM collection_chat_messages WHERE user_id = :user_id"),
            {"user_id": user_id},
        )
        conn.execute(text("DELETE FROM sessions WHERE user_id = :user_id"), {"user_id": user_id})
        conn.execute(text("DELETE FROM users WHERE id = :user_id"), {"user_id": user_id})
    return True


def create_session(user_id):
    token = secrets.token_urlsafe(32)
    with engine.begin() as conn:
        conn.execute(
            sessions.insert().values(
                token=token,
                user_id=user_id,
                created_at=_utc_now_iso(),
            )
        )
    return token


def get_user_by_session_token(token):
    if not token:
        return None
    with engine.begin() as conn:
        row = conn.execute(
            text(
                """
                SELECT user_id
                FROM sessions
                WHERE token = :token
                """
            ),
            {"token": token},
        ).first()
    if not row:
        return None
    return get_user_by_id(row[0])


def login_user(identifier, password):
    normalized_identifier = _normalize_identity(identifier)
    if normalized_identifier == "admin" and password == "admin123":
        admin_user = _admin_user()
        return {**admin_user, "token": create_session(admin_user["id"])}

    with engine.begin() as conn:
        row = conn.execute(
            text(
                """
                SELECT id, username, full_name, email, password_hash, role, is_active, auth_source, created_at
                FROM users
                WHERE lower(email) = :identifier OR lower(username) = :identifier
                """
            ),
            {"identifier": normalized_identifier.lower()},
        ).mappings().first()
    user = dict(row) if row else None
    if user and not user.get("is_active"):
        raise ValueError("This user is inactive.")

    if user and user.get("auth_source") == "local":
        if not _verify_password(password, user["password_hash"]):
            raise ValueError("Invalid credentials.")
    else:
        ad_profile = authenticate_active_directory(normalized_identifier, password)
        if not ad_profile:
            raise ValueError("Invalid credentials.")

        if not user:
            if not ACTIVE_DIRECTORY_AUTO_PROVISION:
                raise ValueError("This Active Directory user has not been provisioned yet.")
            user = create_user(
                full_name=ad_profile["full_name"],
                email=ad_profile["email"],
                password="",
                username=ad_profile["username"],
                role=ACTIVE_DIRECTORY_DEFAULT_ROLE if ACTIVE_DIRECTORY_DEFAULT_ROLE in {"admin", "user"} else "user",
                is_active=1,
                auth_source="active_directory",
            )
        elif user.get("auth_source") != "active_directory":
            raise ValueError("This user must log in with local credentials.")

    user.pop("password_hash", None)
    user["token"] = create_session(user["id"])
    return user


def create_collection(name, document_ids, current_user):
    normalized_name = " ".join((name or "").split())
    if not normalized_name:
        raise ValueError("Collection name is required.")
    if not document_ids:
        raise ValueError("Select at least one document for the collection.")

    valid_documents = []
    for doc_id in dict.fromkeys(document_ids):
        document = get_document(doc_id)
        if not document:
            continue
        if not user_can_access_document(current_user, document):
            raise PermissionError("You cannot add a document you do not own to this collection.")
        valid_documents.append(doc_id)
    if not valid_documents:
        raise ValueError("None of the selected documents were found.")

    collection_id = str(uuid.uuid4())
    created_at = _utc_now_iso()
    with engine.begin() as conn:
        existing = conn.execute(
            text("SELECT id FROM collections WHERE lower(name) = lower(:name)"),
            {"name": normalized_name},
        ).first()
        if existing:
            raise ValueError("A collection with this name already exists.")

        conn.execute(
            collections.insert().values(
                id=collection_id,
                name=normalized_name,
                owner_id=current_user["id"],
                created_at=created_at,
            )
        )
        for doc_id in valid_documents:
            conn.execute(
                collection_documents.insert().values(
                    collection_id=collection_id,
                    doc_id=doc_id,
                )
            )

    return get_collection(collection_id)


def update_collection_documents(collection_id, document_ids, current_user):
    collection = get_collection(collection_id)
    if not collection:
        raise FileNotFoundError("Collection not found.")
    if not user_can_access_collection(current_user, collection):
        raise PermissionError("You cannot edit this collection.")

    valid_documents = []
    for doc_id in dict.fromkeys(document_ids):
        document = get_document(doc_id)
        if not document:
            continue
        if not user_can_access_document(current_user, document):
            raise PermissionError("You cannot add a document you do not own to this collection.")
        valid_documents.append(doc_id)
    if not valid_documents:
        raise ValueError("Select at least one document for the collection.")

    with engine.begin() as conn:
        conn.execute(
            text("DELETE FROM collection_documents WHERE collection_id = :collection_id"),
            {"collection_id": collection_id},
        )
        for doc_id in valid_documents:
            conn.execute(
                collection_documents.insert().values(
                    collection_id=collection_id,
                    doc_id=doc_id,
                )
            )

    return get_collection(collection_id)


def delete_collection(collection_id, current_user):
    collection = get_collection(collection_id)
    if not collection:
        return False
    if not user_can_access_collection(current_user, collection):
        raise PermissionError("You cannot delete this collection.")

    with engine.begin() as conn:
        conn.execute(
            text("DELETE FROM collection_chat_messages WHERE collection_id = :collection_id"),
            {"collection_id": collection_id},
        )
        conn.execute(
            text("DELETE FROM collection_documents WHERE collection_id = :collection_id"),
            {"collection_id": collection_id},
        )
        conn.execute(
            text("DELETE FROM collections WHERE id = :collection_id"),
            {"collection_id": collection_id},
        )
    return True


def find_soffice():
    possible_paths = [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/usr/bin/soffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    for path in possible_paths:
        if os.path.exists(path):
            return path
    raise ValueError("Legacy .doc uploads require LibreOffice to be installed on the server.")


def find_tesseract():
    possible_paths = [
        os.getenv("TESSERACT_CMD", "").strip(),
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        "/usr/bin/tesseract",
        "/opt/homebrew/bin/tesseract",
    ]
    for path in possible_paths:
        if path and os.path.exists(path):
            return path
    return ""


def has_ocr_support():
    return bool(pytesseract and Image and find_tesseract())


def preprocess_ocr_image(image):
    if not (ImageEnhance and ImageFilter and ImageOps):
        return image

    grayscale = ImageOps.grayscale(image)
    autocontrasted = ImageOps.autocontrast(grayscale)
    contrasted = ImageEnhance.Contrast(autocontrasted).enhance(OCR_CONTRAST)
    sharpened = contrasted.filter(ImageFilter.SHARPEN)
    thresholded = sharpened.point(
        lambda pixel: 255 if pixel > OCR_THRESHOLD else 0,
        mode="1",
    )
    return thresholded.convert("L")


def build_ocr_image_variants(image):
    variants = [("original", image.convert("L"))]
    processed = preprocess_ocr_image(image)
    variants.append(("processed", processed))
    variants.append(("autocontrast", ImageOps.autocontrast(image.convert("L"))))
    return variants


def _normalize_whitespace(value):
    return " ".join(str(value).split())


def _block_text_from_pdf(block):
    lines = []
    for line in block.get("lines", []):
        spans = [span.get("text", "") for span in line.get("spans", [])]
        line_text = _normalize_whitespace(" ".join(spans))
        if line_text:
            lines.append(line_text)
    return "\n".join(lines)


def _classify_pdf_block(block, avg_font_size):
    spans = [
        span
        for line in block.get("lines", [])
        for span in line.get("spans", [])
        if span.get("text", "").strip()
    ]
    if not spans:
        return "paragraph"

    font_size = max(span.get("size", 0) for span in spans)
    is_bold = any("bold" in span.get("font", "").lower() or span.get("flags", 0) & 16 for span in spans)
    text_value = _block_text_from_pdf(block)
    if font_size >= avg_font_size * 1.15 or is_bold:
        return "heading"
    if text_value.lstrip().startswith(("-", "*", "\u2022")):
        return "list"

    y_top = block.get("bbox", [0, 0, 0, 0])[1]
    y_bottom = block.get("bbox", [0, 0, 0, 0])[3]
    if font_size <= max(avg_font_size * 0.85, 8) and y_top < 65:
        return "header"
    if font_size <= max(avg_font_size * 0.8, 8) and y_bottom > 700:
        return "footnote"
    return "paragraph"


def extract_structured_text_from_pdf_page(page):
    page_dict = page.get_text("dict")
    text_blocks = [block for block in page_dict.get("blocks", []) if block.get("type") == 0]
    spans = [
        span
        for block in text_blocks
        for line in block.get("lines", [])
        for span in line.get("spans", [])
        if span.get("text", "").strip()
    ]
    avg_font_size = (
        sum(span.get("size", 0) for span in spans) / len(spans)
        if spans
        else 11
    )

    sections = []
    for block in text_blocks:
        block_text = _block_text_from_pdf(block)
        if not block_text:
            continue

        block_type = _classify_pdf_block(block, avg_font_size)
        if block_type == "heading":
            sections.append(f"[Heading]\n{block_text}")
        elif block_type == "header":
            sections.append(f"[Header]\n{block_text}")
        elif block_type == "list":
            normalized_lines = [f"- {line.lstrip('-*• ').strip()}" for line in block_text.splitlines() if line.strip()]
            sections.append("[List]\n" + "\n".join(normalized_lines))
        elif block_type == "footnote":
            sections.append(f"[Footnote]\n{block_text}")
        else:
            sections.append(block_text)

    try:
        table_finder = page.find_tables()
        extracted_tables = table_finder.tables if table_finder else []
    except Exception:
        extracted_tables = []

    for index, table in enumerate(extracted_tables, start=1):
        rows = table.extract() or []
        formatted_rows = []
        for row in rows:
            normalized_cells = [_normalize_whitespace(cell) for cell in row if _normalize_whitespace(cell)]
            if normalized_cells:
                formatted_rows.append(" | ".join(normalized_cells))
        if formatted_rows:
            sections.append(f"[Table {index}]\n" + "\n".join(formatted_rows))

    return "\n\n".join(section for section in sections if section.strip())


def page_requires_ocr(page, extracted_text):
    text_length = len(_normalize_whitespace(extracted_text))
    image_count = len(page.get_images(full=True))
    return text_length < 40 and image_count > 0


def score_ocr_result(text_value, confidence_values):
    normalized = _normalize_whitespace(text_value)
    if not normalized:
        return -1

    alpha_chars = sum(character.isalpha() for character in normalized)
    digit_chars = sum(character.isdigit() for character in normalized)
    useful_confidences = [value for value in confidence_values if value >= 0]
    average_confidence = (
        sum(useful_confidences) / len(useful_confidences) if useful_confidences else 0
    )
    return (average_confidence * 2) + alpha_chars + (digit_chars * 0.3)


def run_ocr_candidate(image, page_segmentation_mode):
    base_config = f"--oem 1 --psm {page_segmentation_mode} preserve_interword_spaces=1"
    data = pytesseract.image_to_data(
        image,
        lang=OCR_LANGUAGE,
        config=base_config,
        output_type=pytesseract.Output.DICT,
    )
    words = [word.strip() for word in data.get("text", []) if word and word.strip()]
    confidences = []
    for value in data.get("conf", []):
        try:
            confidences.append(float(value))
        except (TypeError, ValueError):
            continue
    return " ".join(words), confidences


def extract_ocr_text_from_pdf_page(page):
    tesseract_path = find_tesseract()
    if not (pytesseract and Image and tesseract_path):
        raise ValueError(
            "This PDF appears to be scanned and requires OCR support. Install Tesseract OCR on the server to process scanned PDFs."
        )

    pytesseract.pytesseract.tesseract_cmd = tesseract_path
    pixmap = page.get_pixmap(matrix=fitz.Matrix(OCR_DPI_SCALE, OCR_DPI_SCALE), alpha=False)
    image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
    best_text = ""
    best_score = -1

    for _, variant in build_ocr_image_variants(image):
        for page_segmentation_mode in OCR_PAGE_SEGMENT_MODES:
            try:
                candidate_text, confidences = run_ocr_candidate(variant, page_segmentation_mode)
            except pytesseract.TesseractError:
                continue
            candidate_score = score_ocr_result(candidate_text, confidences)
            if candidate_score > best_score:
                best_text = candidate_text
                best_score = candidate_score

    normalized = _normalize_whitespace(best_text)
    if not normalized:
        raise ValueError("OCR could not extract readable text from the scanned PDF page.")
    return f"[OCR]\n{best_text.strip()}"


def extract_text_from_pdf(path):
    document = fitz.open(path)
    pages = []
    for index in range(document.page_count):
        page = document.load_page(index)
        structured_text = extract_structured_text_from_pdf_page(page)
        if page_requires_ocr(page, structured_text):
            page_text = extract_ocr_text_from_pdf_page(page)
        else:
            page_text = structured_text or page.get_text("text")
        pages.append((index + 1, page_text))
    document.close()
    return pages


def iter_docx_block_items(document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield DocxTable(child, document)


def format_docx_paragraph(paragraph):
    text_value = paragraph.text.strip()
    if not text_value:
        return ""

    style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
    if style_name.startswith("heading"):
        return f"[{paragraph.style.name}]\n{text_value}"
    if "list" in style_name or text_value.startswith(("-", "*", "\u2022")):
        return f"[List]\n- {text_value.lstrip('-*• ').strip()}"
    if "footer" in style_name:
        return f"[Footer]\n{text_value}"
    if "header" in style_name:
        return f"[Header]\n{text_value}"
    if "footnote" in style_name:
        return f"[Footnote]\n{text_value}"
    return text_value


def format_docx_table(table, table_index):
    rows = []
    for row in table.rows:
        cells = [_normalize_whitespace(cell.text) for cell in row.cells if _normalize_whitespace(cell.text)]
        if cells:
            rows.append(" | ".join(cells))
    if not rows:
        return ""
    return f"[Table {table_index}]\n" + "\n".join(rows)


def extract_docx_header_footer_sections(document):
    sections = []
    table_index = 1
    for section_index, section in enumerate(document.sections, start=1):
        for label, container in (
            (f"Section {section_index} Header", section.header),
            (f"Section {section_index} Footer", section.footer),
        ):
            block_texts = []
            for paragraph in container.paragraphs:
                text_value = paragraph.text.strip()
                if text_value:
                    block_texts.append(text_value)
            for table in container.tables:
                formatted_table = format_docx_table(table, table_index)
                if formatted_table:
                    block_texts.append(formatted_table)
                    table_index += 1
            if block_texts:
                sections.append(f"[{label}]\n" + "\n\n".join(block_texts))
    return sections


def extract_text_from_docx(path):
    document = docx.Document(path)
    sections = []
    table_index = 1

    for block in iter_docx_block_items(document):
        if isinstance(block, Paragraph):
            formatted = format_docx_paragraph(block)
        else:
            formatted = format_docx_table(block, table_index)
            if formatted:
                table_index += 1
        if formatted:
            sections.append(formatted)

    sections.extend(extract_docx_header_footer_sections(document))

    text_content = "\n\n".join(sections)
    return [(1, text_content)]


def extract_text_from_doc(path):
    soffice_path = find_soffice()
    with tempfile.TemporaryDirectory() as temp_dir:
        try:
            subprocess.run(
                [soffice_path, "--headless", "--convert-to", "docx", path, "--outdir", temp_dir],
                check=True,
                capture_output=True,
                timeout=90,
            )
        except subprocess.CalledProcessError as exc:
            stderr_output = (exc.stderr or b"").decode("utf-8", errors="ignore").strip()
            raise ValueError(
                "This .doc file could not be converted to a readable format."
                + (f" LibreOffice said: {stderr_output}" if stderr_output else "")
            ) from exc
        except subprocess.TimeoutExpired as exc:
            raise ValueError("This .doc file took too long to convert.") from exc
        converted_files = [name for name in os.listdir(temp_dir) if name.lower().endswith(".docx")]
        if not converted_files:
            raise ValueError("LibreOffice did not produce a DOCX file.")
        return extract_text_from_docx(os.path.join(temp_dir, converted_files[0]))


def convert_legacy_spreadsheet_to_xlsx(path):
    soffice_path = find_soffice()
    with tempfile.TemporaryDirectory() as temp_dir:
        try:
            subprocess.run(
                [soffice_path, "--headless", "--convert-to", "xlsx", path, "--outdir", temp_dir],
                check=True,
                capture_output=True,
                timeout=120,
            )
        except subprocess.CalledProcessError as exc:
            stderr_output = (exc.stderr or b"").decode("utf-8", errors="ignore").strip()
            raise ValueError(
                "This .xls file could not be converted to a formula-aware workbook."
                + (f" LibreOffice said: {stderr_output}" if stderr_output else "")
            ) from exc
        except subprocess.TimeoutExpired as exc:
            raise ValueError("This .xls file took too long to convert.") from exc

        converted_files = [name for name in os.listdir(temp_dir) if name.lower().endswith(".xlsx")]
        if not converted_files:
            raise ValueError("LibreOffice did not produce an XLSX file.")

        return extract_text_from_xlsx(os.path.join(temp_dir, converted_files[0]))


def extract_text_from_xlsx(path):
    workbook = load_workbook(path, data_only=False)
    pages = []
    for sheet_index, sheet_name in enumerate(workbook.sheetnames, start=1):
        sheet = workbook[sheet_name]
        rows = []
        formula_count = 0
        merged_ranges = [str(cell_range) for cell_range in sheet.merged_cells.ranges]
        for row in sheet.iter_rows():
            populated_cells = []
            coords = []
            for cell in row:
                value = cell.value
                if value in (None, ""):
                    continue
                coords.append(cell.coordinate)
                if isinstance(value, str) and value.startswith("="):
                    formula_count += 1
                    populated_cells.append(f"{cell.coordinate}: formula {value}")
                else:
                    populated_cells.append(f"{cell.coordinate}: {value}")
            if populated_cells:
                row_range = f"{coords[0]}:{coords[-1]}" if len(coords) > 1 else coords[0]
                rows.append(f"[Range {row_range}] " + " | ".join(populated_cells))

        sheet_meta = (
            f"[Sheet: {sheet_name}] rows={sheet.max_row}, cols={sheet.max_column}, "
            f"formulas={formula_count}, merged_ranges={len(merged_ranges)}, "
            f"frozen_panes={sheet.freeze_panes or 'none'}"
        )
        merged_meta = f"[Merged Ranges] {', '.join(merged_ranges)}" if merged_ranges else ""
        content_parts = [sheet_meta]
        if merged_meta:
            content_parts.append(merged_meta)
        content_parts.extend(rows)
        pages.append((sheet_index, "\n".join(content_parts)))
    return pages


def extract_text_from_xls(path):
    try:
        return convert_legacy_spreadsheet_to_xlsx(path)
    except ValueError:
        pass

    try:
        workbook = xlrd.open_workbook(path, formatting_info=False)
    except Exception as exc:
        raise ValueError(
            "This .xls file could not be opened. It may be corrupted or not a real legacy Excel file."
        ) from exc

    pages = []
    for sheet_index in range(workbook.nsheets):
        sheet = workbook.sheet_by_index(sheet_index)
        rows = []
        for row_index in range(sheet.nrows):
            populated_cells = []
            coords = []
            for col_index in range(sheet.ncols):
                value = sheet.cell_value(row_index, col_index)
                if value in ("", None):
                    continue
                coord = f"{get_column_letter(col_index + 1)}{row_index + 1}"
                coords.append(coord)
                populated_cells.append(f"{coord}: {value}")
            if populated_cells:
                row_range = f"{coords[0]}:{coords[-1]}" if len(coords) > 1 else coords[0]
                rows.append(f"[Range {row_range}] " + " | ".join(populated_cells))

        sheet_meta = (
            f"[Sheet: {sheet.name}] rows={sheet.nrows}, cols={sheet.ncols}, formulas=legacy-values-only"
        )
        pages.append((sheet_index + 1, sheet_meta + "\n" + "\n".join(rows)))
    return pages


def extract_pages(path):
    extension = Path(path).suffix.lower().lstrip(".")
    if extension == "pdf":
        return extract_text_from_pdf(path)
    if extension == "docx":
        return extract_text_from_docx(path)
    if extension == "doc":
        return extract_text_from_doc(path)
    if extension == "xlsx":
        return extract_text_from_xlsx(path)
    if extension == "xls":
        return extract_text_from_xls(path)
    raise ValueError(f"Unsupported file type: {extension}")


def chunk_text(page_text, page_num, chunk_size=CHUNK_SIZE_CHARS, overlap=CHUNK_OVERLAP):
    cleaned = page_text.strip()
    if not cleaned:
        return []

    chunk_records = []
    start = 0
    chunk_index = 0
    while start < len(cleaned):
        end = min(len(cleaned), start + chunk_size)
        chunk_records.append(
            {
                "page": page_num,
                "text": cleaned[start:end],
                "char_start": start,
                "char_end": end,
                "chunk_index": chunk_index,
            }
        )
        chunk_index += 1
        if end >= len(cleaned):
            break
        start = max(end - overlap, 0)
    return chunk_records


def _tag_from_block(block_text):
    stripped = block_text.strip()
    if stripped.startswith("[") and "]" in stripped:
        return stripped[1:stripped.index("]")].strip()
    return ""


def _strip_tag_prefix(block_text):
    stripped = block_text.strip()
    if stripped.startswith("[") and "]" in stripped:
        return stripped[stripped.index("]") + 1 :].strip()
    return stripped


def build_document_chunk_records(pages, file_type):
    chunk_records = []

    if file_type in {"xlsx", "xls"}:
        for page_num, page_text in pages:
            lines = [line.strip() for line in page_text.splitlines() if line.strip()]
            current_sheet = ""
            sheet_meta = ""
            paragraph_index = 0

            for line in lines:
                if line.startswith("[Sheet:"):
                    current_sheet = line.split("]", 1)[0].replace("[Sheet:", "").strip()
                    sheet_meta = line
                    continue

                if line.startswith("[Merged Ranges]"):
                    if chunk_records and chunk_records[-1].get("sheet_name") == current_sheet:
                        chunk_records[-1]["text"] += f"\n{line}"
                    continue

                cell_range = ""
                source_ref = current_sheet or f"Sheet {page_num}"
                if line.startswith("[Range ") and "]" in line:
                    cell_range = line.split("]", 1)[0].replace("[Range", "").strip()
                    source_ref = f"{current_sheet}!{cell_range}" if current_sheet else cell_range

                paragraph_index += 1
                chunk_records.append(
                    {
                        "page": page_num,
                        "text": f"{sheet_meta}\n{line}" if sheet_meta else line,
                        "char_start": 0,
                        "char_end": len(line),
                        "section_label": "sheet-row",
                        "sheet_name": current_sheet or None,
                        "cell_range": cell_range or None,
                        "source_ref": source_ref,
                        "chunk_index": len(chunk_records),
                    }
                )

        return chunk_records

    for page_num, page_text in pages:
        blocks = [block.strip() for block in page_text.split("\n\n") if block.strip()]
        paragraph_index = 0

        for block in blocks:
            tag = _tag_from_block(block)
            body = _strip_tag_prefix(block)
            if not body:
                continue

            section_label = tag.lower() if tag else "paragraph"
            if tag.lower().startswith("table"):
                source_ref = f"Page {page_num} {tag}"
            elif tag.lower().startswith("section"):
                source_ref = tag
            else:
                paragraph_index += 1
                source_ref = f"Page {page_num} Paragraph {paragraph_index}"

            chunk_records.append(
                {
                    "page": page_num,
                    "text": block,
                    "char_start": 0,
                    "char_end": len(body),
                    "section_label": section_label,
                    "sheet_name": None,
                    "cell_range": None,
                    "source_ref": source_ref,
                    "chunk_index": len(chunk_records),
                }
            )

    return chunk_records


def _persist_document(doc_id, filename, stored_path, file_hash, file_size, file_type, pages, chunk_records, owner_id):
    with engine.begin() as conn:
        conn.execute(
            documents.insert().values(
                id=doc_id,
                filename=filename,
                stored_path=str(stored_path),
                file_hash=file_hash,
                file_size=file_size,
                file_type=file_type,
                page_count=len(pages),
                chunk_count=len(chunk_records),
                owner_id=owner_id,
                created_at=_utc_now_iso(),
            )
        )
        for record in chunk_records:
            conn.execute(
                chunks.insert().values(
                    id=record["id"],
                    doc_id=doc_id,
                chunk_index=record["chunk_index"],
                text=record["text"],
                page=record["page"],
                char_start=record["char_start"],
                char_end=record["char_end"],
                section_label=record.get("section_label"),
                sheet_name=record.get("sheet_name"),
                cell_range=record.get("cell_range"),
                source_ref=record.get("source_ref"),
            )
        )


def ingest_file(filepath, original_filename=None, owner_id="admin"):
    # Keep parsing, file storage, chunking, embeddings, and metadata writes inside one flow for a single upload action.
    source_path = Path(filepath)
    validation = validate_upload_file(source_path, original_filename=original_filename)

    file_hash = compute_file_hash(source_path)
    existing = get_document_by_hash(file_hash, owner_id)
    if existing:
        return existing

    filename = validation["filename"]
    extension = validation["extension"]
    file_size = validation["file_size"]

    doc_id = str(uuid.uuid4())
    storage_filename = f"{doc_id}_{filename}"
    local_storage_path = DATA_DIR / storage_filename
    remote_storage_path = None

    # Local mode keeps a working copy in the project folder; Nextcloud mode indexes first, then uploads remotely.
    if is_nextcloud_enabled():
        working_path = source_path
    else:
        shutil.copy2(source_path, local_storage_path)
        working_path = local_storage_path

    try:
        pages = extract_pages(str(working_path))
        chunk_records = build_document_chunk_records(pages, extension)
        chunk_texts = []

        for record in chunk_records:
            record["id"] = str(uuid.uuid4())
            record["doc_id"] = doc_id
            record["chunk_index"] = len(chunk_texts)
            chunk_texts.append(record["text"])

        if not chunk_texts:
            raise ValueError("No text content could be extracted from the file.")

        embeddings = get_embedder().encode(
            chunk_texts,
            convert_to_numpy=True,
            show_progress_bar=False,
        ).astype("float32")

        index = faiss.IndexFlatL2(embeddings.shape[1])
        index.add(embeddings)
        faiss.write_index(index, str(_index_path(doc_id)))
        np.save(_mapping_path(doc_id), np.array([record["id"] for record in chunk_records], dtype=object))

        # Persist the original file in Nextcloud when configured, otherwise keep the local working copy as the stored file.
        if is_nextcloud_enabled():
            remote_storage_path = upload_file(str(source_path), storage_filename)
            stored_path = remote_storage_path
        else:
            stored_path = local_storage_path

        _persist_document(
            doc_id=doc_id,
            filename=filename,
            stored_path=stored_path,
            file_hash=file_hash,
            file_size=file_size,
            file_type=extension,
            pages=pages,
            chunk_records=chunk_records,
            owner_id=owner_id,
        )
    except Exception:
        # Roll back whichever storage target was used if parsing or indexing fails midway through ingestion.
        if remote_storage_path:
            try:
                delete_file(remote_storage_path)
            except Exception:
                pass
        if local_storage_path.exists():
            local_storage_path.unlink()
        if _index_path(doc_id).exists():
            _index_path(doc_id).unlink()
        if _mapping_path(doc_id).exists():
            _mapping_path(doc_id).unlink()
        raise

    return get_document(doc_id)


def delete_document(doc_id):
    document = get_document(doc_id)
    if not document:
        return False

    with engine.begin() as conn:
        conn.execute(text("DELETE FROM chat_messages WHERE doc_id = :doc_id"), {"doc_id": doc_id})
        conn.execute(text("DELETE FROM collection_documents WHERE doc_id = :doc_id"), {"doc_id": doc_id})
        conn.execute(text("DELETE FROM chunks WHERE doc_id = :doc_id"), {"doc_id": doc_id})
        conn.execute(text("DELETE FROM documents WHERE id = :doc_id"), {"doc_id": doc_id})

    stored_path = document["stored_path"]
    if stored_path.startswith("http://") or stored_path.startswith("https://"):
        delete_file(stored_path)
    else:
        local_path = Path(stored_path)
        if local_path.exists():
            local_path.unlink()

    for path in [_index_path(doc_id), _mapping_path(doc_id)]:
        if path.exists():
            path.unlink()

    return True


def save_chat_message(doc_id, role, message, user_id, model=None, citations_json=None):
    with engine.begin() as conn:
        conn.execute(
            chat_messages.insert().values(
                id=str(uuid.uuid4()),
                doc_id=doc_id,
                role=role,
                message=message,
                user_id=user_id,
                model=model,
                citations_json=citations_json,
                created_at=_utc_now_iso(),
            )
        )


def save_collection_chat_message(collection_id, role, message, user_id, model=None, citations_json=None):
    with engine.begin() as conn:
        conn.execute(
            collection_chat_messages.insert().values(
                id=str(uuid.uuid4()),
                collection_id=collection_id,
                role=role,
                message=message,
                user_id=user_id,
                model=model,
                citations_json=citations_json,
                created_at=_utc_now_iso(),
            )
        )


def list_chat_messages(doc_id, current_user):
    params = {"doc_id": doc_id}
    if current_user.get("role") == "admin":
        where_clause = "WHERE doc_id = :doc_id"
    else:
        where_clause = "WHERE doc_id = :doc_id AND user_id = :user_id"
        params["user_id"] = current_user["id"]
    with engine.begin() as conn:
        rows = conn.execute(
            text(
                """
                SELECT id, doc_id, role, message, user_id, model, citations_json, created_at
                FROM chat_messages
                {where_clause}
                ORDER BY created_at ASC
                """
                .format(where_clause=where_clause)
            ),
            params,
        ).mappings().all()
    return [dict(row) for row in rows]


def list_collection_chat_messages(collection_id, current_user):
    params = {"collection_id": collection_id}
    if current_user.get("role") == "admin":
        where_clause = "WHERE collection_id = :collection_id"
    else:
        where_clause = "WHERE collection_id = :collection_id AND user_id = :user_id"
        params["user_id"] = current_user["id"]
    with engine.begin() as conn:
        rows = conn.execute(
            text(
                """
                SELECT id, collection_id, role, message, user_id, model, citations_json, created_at
                FROM collection_chat_messages
                {where_clause}
                ORDER BY created_at ASC
                """
                .format(where_clause=where_clause)
            ),
            params,
        ).mappings().all()
    return [dict(row) for row in rows]


def clear_chat_messages(doc_id, current_user):
    params = {"doc_id": doc_id}
    if current_user.get("role") == "admin":
        query = "DELETE FROM chat_messages WHERE doc_id = :doc_id"
    else:
        query = "DELETE FROM chat_messages WHERE doc_id = :doc_id AND user_id = :user_id"
        params["user_id"] = current_user["id"]
    with engine.begin() as conn:
        result = conn.execute(
            text(query),
            params,
        )
    return result.rowcount or 0


def clear_collection_chat_messages(collection_id, current_user):
    params = {"collection_id": collection_id}
    if current_user.get("role") == "admin":
        query = "DELETE FROM collection_chat_messages WHERE collection_id = :collection_id"
    else:
        query = (
            "DELETE FROM collection_chat_messages "
            "WHERE collection_id = :collection_id AND user_id = :user_id"
        )
        params["user_id"] = current_user["id"]
    with engine.begin() as conn:
        result = conn.execute(
            text(query),
            params,
        )
    return result.rowcount or 0


def delete_chat_message(doc_id, message_id, current_user):
    params = {"doc_id": doc_id, "message_id": message_id}
    if current_user.get("role") == "admin":
        query = """
                DELETE FROM chat_messages
                WHERE doc_id = :doc_id AND id = :message_id
                """
    else:
        query = """
                DELETE FROM chat_messages
                WHERE doc_id = :doc_id AND id = :message_id AND user_id = :user_id
                """
        params["user_id"] = current_user["id"]
    with engine.begin() as conn:
        result = conn.execute(
            text(query),
            params,
        )
    return result.rowcount or 0


def delete_collection_chat_message(collection_id, message_id, current_user):
    params = {"collection_id": collection_id, "message_id": message_id}
    if current_user.get("role") == "admin":
        query = """
                DELETE FROM collection_chat_messages
                WHERE collection_id = :collection_id AND id = :message_id
                """
    else:
        query = """
                DELETE FROM collection_chat_messages
                WHERE collection_id = :collection_id AND id = :message_id AND user_id = :user_id
                """
        params["user_id"] = current_user["id"]
    with engine.begin() as conn:
        result = conn.execute(
            text(query),
            params,
        )
    return result.rowcount or 0
